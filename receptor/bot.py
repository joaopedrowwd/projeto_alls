import os
import sqlite3
import threading
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing

import requests
from docx import Document
from dotenv import load_dotenv
from flask import Flask, request
from langchain_groq import ChatGroq
from pydantic import BaseModel, Field

# =====================================================================
# CONFIGURAÇÃO
# =====================================================================
load_dotenv()
GROQ_API_KEY = os.getenv("GROQ_API_KEY")

model = ChatGroq(model="openai/gpt-oss-120b", api_key=GROQ_API_KEY, temperature=0.2)

# Evolution API
API = os.getenv("EVOLUTION_API", "http://localhost:8080")
KEY = os.getenv("EVOLUTION_KEY", "projeto_alls")
INSTANCIA = os.getenv("EVOLUTION_INSTANCIA", "Teste")

# Banco e contexto
DB = "bot.db"
MAX_NAO_RESUMIDAS = 12   # acima disso, as mensagens antigas viram resumo
MANTER_ULTIMAS = 6       # quantas mensagens recentes ficam na íntegra

# Regras de negócio (arquivo Word na pasta do projeto)
ARQUIVO_REGRAS = "regras_negocio.docx"

INSTRUCOES_FIXAS = (
    "Você só responderá mensagens relacionadas à empresa de canetas, desde regra de negócio. "
    "Sempre ao explicar para o cliente, na hora que ele se interessar mais, "
    "ofereça o acesso ao site https://projetoalls.com.br (envie sempre a URL pura assim, "
    "NUNCA use formato Markdown como [texto](url), pois você está no WhatsApp). "
    "Para negrito no WhatsApp use apenas um asterisco (*texto*). "
    "Se ainda não souber o nome do cliente, pergunte de forma natural em algum momento. "
    "Responda apenas com base nas regras abaixo. Se a informação não estiver nelas, "
    "diga que vai verificar com a equipe, sem inventar preços, prazos ou políticas."
)
app = Flask(__name__)
pool = ThreadPoolExecutor(max_workers=8)
locks = defaultdict(threading.Lock)   # um lock por cliente (evita mensagens simultâneas)


# =====================================================================
# REGRAS DE NEGÓCIO (WORD)
# =====================================================================
_cache_regras = {"mtime": None, "texto": ""}


def ler_docx(caminho):
    doc = Document(caminho)
    partes = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for tabela in doc.tables:  # tabelas (preços, prazos etc.)
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells]
            partes.append(" | ".join(celulas))
    return "\n".join(partes)


def carregar_regras():
    """Lê o Word só quando o arquivo muda; senão usa o cache."""
    try:
        mtime = os.path.getmtime(ARQUIVO_REGRAS)
    except FileNotFoundError:
        print(f"Aviso: {ARQUIVO_REGRAS} não encontrado.")
        return ""
    if mtime != _cache_regras["mtime"]:
        _cache_regras["texto"] = ler_docx(ARQUIVO_REGRAS)
        _cache_regras["mtime"] = mtime
        print(f"Regras recarregadas (~{len(_cache_regras['texto']) // 3} tokens).")
    return _cache_regras["texto"]


def system_base():
    regras = carregar_regras()
    if not regras:
        return INSTRUCOES_FIXAS
    return f"{INSTRUCOES_FIXAS}\n\nREGRAS DE NEGÓCIO DA EMPRESA:\n{regras}"


# =====================================================================
# BANCO DE DADOS (SQLite)
# =====================================================================
def db():
    con = sqlite3.connect(DB, timeout=10)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL")
    return con


def init_db():
    with closing(db()) as con, con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS clientes(
                jid TEXT PRIMARY KEY,
                nome TEXT,               -- nome que o cliente informou
                nome_whatsapp TEXT,      -- pushName (não confirmado)
                fatos TEXT DEFAULT '',   -- preferências, interesses, pendências
                resumo TEXT DEFAULT '',
                atualizado_em TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE IF NOT EXISTS mensagens(
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                jid TEXT NOT NULL,
                role TEXT NOT NULL,
                texto TEXT NOT NULL,
                resumida INTEGER DEFAULT 0,
                ts TEXT DEFAULT CURRENT_TIMESTAMP
            );
            CREATE INDEX IF NOT EXISTS idx_msg ON mensagens(jid, id);
            CREATE TABLE IF NOT EXISTS processadas(msg_id TEXT PRIMARY KEY);
            """
        )


def get_cliente(jid, push_name):
    with closing(db()) as con, con:
        con.execute(
            "INSERT INTO clientes(jid, nome_whatsapp) VALUES(?,?) "
            "ON CONFLICT(jid) DO UPDATE SET "
            "nome_whatsapp=COALESCE(excluded.nome_whatsapp, nome_whatsapp)",
            (jid, push_name),
        )
        return dict(con.execute("SELECT * FROM clientes WHERE jid=?", (jid,)).fetchone())


def salvar_msg(jid, role, texto):
    with closing(db()) as con, con:
        con.execute("INSERT INTO mensagens(jid, role, texto) VALUES(?,?,?)", (jid, role, texto))


def nao_resumidas(jid):
    with closing(db()) as con:
        rows = con.execute(
            "SELECT id, role, texto FROM mensagens WHERE jid=? AND resumida=0 ORDER BY id",
            (jid,),
        ).fetchall()
    return [dict(r) for r in rows]


def ja_processada(msg_id):
    """Evita responder duas vezes se o webhook for reenviado."""
    with closing(db()) as con, con:
        try:
            con.execute("INSERT INTO processadas(msg_id) VALUES(?)", (msg_id,))
            return False
        except sqlite3.IntegrityError:
            return True


# =====================================================================
# EXTRAÇÃO DE NOME E RESUMO
# =====================================================================
class Extracao(BaseModel):
    nome: str | None = Field(
        None, description="Nome que o cliente disse ser o dele. None se não disse."
    )


class ResumoOut(BaseModel):
    resumo: str = Field(description="Resumo da conversa em até 120 palavras")
    fatos: list[str] = Field(
        description="Fatos duráveis do cliente: interesses, preferências, pendências. Máx 8."
    )


def extrair_nome(ultima_ia, texto):
    try:
        out = model.with_structured_output(Extracao).invoke(
            f"Última mensagem do atendente: {ultima_ia or '(início)'}\n"
            f"Mensagem do cliente: {texto}\n"
            "O cliente informou o PRÓPRIO nome? Se sim, retorne só o nome; senão, None."
        )
        return out.nome.strip() if out and out.nome else None
    except Exception as e:
        print("Erro extraindo nome:", e)
        return None


def resumir(jid, cliente):
    msgs = nao_resumidas(jid)
    if len(msgs) <= MAX_NAO_RESUMIDAS:
        return
    antigas = msgs[:-MANTER_ULTIMAS]
    conversa = "\n".join(f"{m['role']}: {m['texto']}" for m in antigas)
    try:
        out = model.with_structured_output(ResumoOut).invoke(
            "Atualize o resumo e os fatos do cliente com as novas mensagens. "
            "Preserve números, nomes, decisões e pendências; descarte saudações e repetição.\n\n"
            f"RESUMO ATUAL: {cliente['resumo'] or '(vazio)'}\n"
            f"FATOS ATUAIS: {cliente['fatos'] or '(vazio)'}\n\n"
            f"NOVAS MENSAGENS:\n{conversa}"
        )
    except Exception as e:
        print("Erro resumindo:", e)
        return
    with closing(db()) as con, con:
        con.execute(
            "UPDATE clientes SET resumo=?, fatos=?, atualizado_em=CURRENT_TIMESTAMP WHERE jid=?",
            (out.resumo, "\n".join(out.fatos), jid),
        )
        con.execute(
            "UPDATE mensagens SET resumida=1 WHERE jid=? AND id<=?",
            (jid, antigas[-1]["id"]),
        )


# =====================================================================
# MONTAGEM DO CONTEXTO
# =====================================================================
def montar(cliente, historico, texto):
    system = system_base()

    ficha = []
    if cliente["nome"]:
        ficha.append(f"Nome do cliente: {cliente['nome']}")
    elif cliente["nome_whatsapp"]:
        ficha.append(f"Nome no WhatsApp (não confirmado): {cliente['nome_whatsapp']}")
    if cliente["fatos"]:
        ficha.append(f"Fatos conhecidos:\n{cliente['fatos']}")
    if cliente["resumo"]:
        ficha.append(f"Resumo da conversa anterior:\n{cliente['resumo']}")
    if ficha:
        system += "\n\n" + "\n".join(ficha)

    msgs = [("system", system)]
    msgs += [("human" if m["role"] == "human" else "ai", m["texto"]) for m in historico]
    msgs.append(("human", texto))
    return msgs


# =====================================================================
# FLUXO PRINCIPAL
# =====================================================================
def enviar(destino, texto):
    r = requests.post(
        f"{API}/message/sendText/{INSTANCIA}",
        headers={"apikey": KEY},
        json={"number": destino, "text": texto},
        timeout=15,
    )
    if not r.ok:
        print("Erro ao enviar:", r.status_code, r.text)


def processar(jid, texto, msg_id, push_name):
    with locks[jid]:
        if msg_id and ja_processada(msg_id):
            return

        cliente = get_cliente(jid, push_name)
        historico = nao_resumidas(jid)

        # Só tenta extrair o nome enquanto ele ainda não é conhecido
        if not cliente["nome"]:
            ultima_ia = next(
                (m["texto"] for m in reversed(historico) if m["role"] == "ai"), None
            )
            nome = extrair_nome(ultima_ia, texto)
            if nome:
                with closing(db()) as con, con:
                    con.execute("UPDATE clientes SET nome=? WHERE jid=?", (nome, jid))
                cliente["nome"] = nome

        try:
            resposta = model.invoke(montar(cliente, historico, texto)).content
        except Exception as e:
            print("Erro no modelo:", e)
            return

        # Só salva no histórico depois que o modelo respondeu com sucesso
        salvar_msg(jid, "human", texto)
        salvar_msg(jid, "ai", resposta)
        enviar(jid, resposta)
        resumir(jid, get_cliente(jid, None))


@app.post("/webhook")
def webhook():
    body = request.get_json(silent=True) or {}
    data = body.get("data") or {}
    key = data.get("key") or {}

    if body.get("event") != "messages.upsert" or key.get("fromMe"):
        return "", 200

    jid = key.get("remoteJid", "")
    if not jid or jid.endswith("@g.us"):  # ignora grupos
        return "", 200

    message = data.get("message") or {}
    texto = message.get("conversation") or (
        message.get("extendedTextMessage") or {}
    ).get("text")
    if not texto:
        return "", 200

    print(f"Recebida de {jid}: {texto}")

    # Responde 200 na hora e processa em segundo plano
    pool.submit(processar, jid, texto.strip()[:4000], key.get("id"), data.get("pushName"))
    return "", 200


if __name__ == "__main__":
    init_db()
    carregar_regras()
    app.run(host="0.0.0.0", port=3000)